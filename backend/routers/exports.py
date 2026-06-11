from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from database import get_db
from models import Client, Vente, Formation, StockEscargot, StockHanneton
from auth import get_current_user
import io
import openpyxl
from openpyxl.styles import Font, PatternFill

router = APIRouter()

GREEN = "16A34A"
WHITE = "FFFFFF"


def header_style(cell):
    cell.font = Font(bold=True, color=WHITE)
    cell.fill = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")


def make_excel(headers, rows, sheet_name="Export"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        header_style(cell)
    for row_idx, row in enumerate(rows, 2):
        for col_idx, val in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


@router.get("/clients/excel")
def export_clients_excel(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    clients = db.query(Client).all()
    headers = ["ID", "Nom", "Prénom", "Téléphone", "WhatsApp", "Ville", "Commune", "Pays", "Profession", "Source", "Statut", "Date d'ajout"]
    rows = [
        (c.id, c.nom, c.prenom, c.telephone, c.whatsapp, c.ville, c.commune,
         c.pays, c.profession, c.source, c.statut,
         c.date_ajout.strftime("%d/%m/%Y") if c.date_ajout else "")
        for c in clients
    ]
    stream = make_excel(headers, rows, "Clients")
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=clients.xlsx"}
    )


@router.get("/ventes/excel")
def export_ventes_excel(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    ventes = db.query(Vente).all()
    headers = ["ID", "Client", "Produit", "Quantité", "Prix Unitaire", "Montant Total", "Montant Payé", "Reste", "Date"]
    rows = [
        (v.id,
         f"{v.client.nom} {v.client.prenom or ''}" if v.client else "",
         v.produit, v.quantite, v.prix_unitaire, v.montant_total,
         v.montant_paye, v.reste_a_payer,
         v.date.strftime("%d/%m/%Y") if v.date else "")
        for v in ventes
    ]
    stream = make_excel(headers, rows, "Ventes")
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=ventes.xlsx"}
    )


@router.get("/formations/excel")
def export_formations_excel(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    formations = db.query(Formation).all()
    headers = ["ID", "Nom", "Date", "Lieu", "Prix", "Nb Participants"]
    rows = [
        (f.id, f.nom, str(f.date), f.lieu, f.prix, len(f.participants))
        for f in formations
    ]
    stream = make_excel(headers, rows, "Formations")
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=formations.xlsx"}
    )


@router.get("/stock/excel")
def export_stock_excel(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    wb = openpyxl.Workbook()

    ws_e = wb.active
    ws_e.title = "Escargots"
    headers_e = ["ID", "Date", "Reproducteurs", "Juvéniles", "Naissances", "Mortalité", "Notes"]
    for col, h in enumerate(headers_e, 1):
        header_style(ws_e.cell(row=1, column=col, value=h))
    for row, s in enumerate(db.query(StockEscargot).all(), 2):
        for col, val in enumerate([s.id, str(s.date), s.reproducteurs, s.juveniles, s.naissances, s.mortalite, s.notes or ""], 1):
            ws_e.cell(row=row, column=col, value=val)

    ws_h = wb.create_sheet("Hannetons")
    headers_h = ["ID", "Date", "Stock Actuel", "Production", "Mortalité", "Notes"]
    for col, h in enumerate(headers_h, 1):
        header_style(ws_h.cell(row=1, column=col, value=h))
    for row, s in enumerate(db.query(StockHanneton).all(), 2):
        for col, val in enumerate([s.id, str(s.date), s.stock_actuel, s.production, s.mortalite, s.notes or ""], 1):
            ws_h.cell(row=row, column=col, value=val)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=stock.xlsx"}
    )
